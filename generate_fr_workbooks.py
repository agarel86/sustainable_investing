"""Create French copies of chapter question workbooks (EN -> FR)."""
from pathlib import Path
import time
import shutil
import openpyxl

try:
    from deep_translator import GoogleTranslator
except ImportError:
    GoogleTranslator = None

BASE = Path(__file__).parent
TRANSLATOR = None
CHUNK = 4500


def tr(text):
    if not text or not str(text).strip():
        return text
    s = str(text)
    global TRANSLATOR
    if TRANSLATOR is None:
        if GoogleTranslator is None:
            raise RuntimeError("Install deep-translator: pip install deep-translator")
        TRANSLATOR = GoogleTranslator(source="en", target="fr")
    out = []
    for i in range(0, len(s), CHUNK):
        part = s[i : i + CHUNK]
        for attempt in range(3):
            try:
                out.append(TRANSLATOR.translate(part))
                break
            except Exception:
                time.sleep(1.5 * (attempt + 1))
        else:
            out.append(part)
        time.sleep(0.15)
    return "".join(out)


def translate_workbook(src: Path, dst: Path):
    wb = openpyxl.load_workbook(src)
    if "All Questions" in wb.sheetnames:
        ws = wb["All Questions"]
        for row in ws.iter_rows(min_row=2):
            if row[2].value:
                row[2].value = tr(row[2].value)
            if len(row) > 3 and row[3].value:
                row[3].value = tr(row[3].value)
    if "Numerical Choices" in wb.sheetnames:
        ws2 = wb["Numerical Choices"]
        for row in ws2.iter_rows(min_row=2):
            if row[1].value:
                row[1].value = tr(row[1].value)
            if row[2].value:
                row[2].value = tr(row[2].value)
            for c in range(3, 8):
                if c < len(row) and row[c].value:
                    row[c].value = tr(str(row[c].value))
    wb.save(dst)
    print(f"  -> {dst.name}")


def main():
    for src in sorted(BASE.glob("Chapter_*_Questions.xlsx")):
        if src.name.endswith("_FR.xlsx"):
            continue
        dst = src.with_name(src.stem + "_FR.xlsx")
        print(f"Translating {src.name}...")
        translate_workbook(src, dst)


if __name__ == "__main__":
    main()
