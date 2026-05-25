import openpyxl
from pathlib import Path

base = Path(__file__).parent
lines = []
for fname in sorted(base.glob("Chapter_*_Questions.xlsx")):
    if "_FR" in fname.name:
        continue
    wb = openpyxl.load_workbook(fname, read_only=True, data_only=True)
    if "Numerical Choices" not in wb.sheetnames:
        wb.close()
        continue
    ws = wb["All Questions"]
    ws2 = wb["Numerical Choices"]
    lines.append(f"=== {fname.name} ===")
    num_rows = {r[0]: r for r in ws2.iter_rows(min_row=2, values_only=True) if r and r[0]}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or (row[1] or "").strip() != "Numerical":
            continue
        qid = row[0]
        lines.append(f"ID {qid}")
        lines.append(f"Q: {row[2]}")
        lines.append(f"A: {row[3]}")
        nr = num_rows.get(qid)
        if nr:
            hdr_note = "cols 3-7"
            for i in range(3, 8):
                if i < len(nr):
                    lines.append(f"  ch{i-2}: {nr[i]}")
            lines.append(f"  correct: {nr[8] if len(nr)>8 else ''}")
        lines.append("")
    wb.close()

(base / "_nums.txt").write_text("\n".join(lines), encoding="utf-8")
print("wrote", len(lines), "lines")
